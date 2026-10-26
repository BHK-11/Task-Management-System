import io
import unittest
import zipfile

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen import canvas

from document_extractor import extract_document


class DocumentExtractorTest(unittest.TestCase):
    def test_extracts_plain_text_and_markdown_lines(self):
        self.assertEqual(extract_document("txt", b"First line\nSecond line")["characterCount"], 21)
        markdown = extract_document("md", "# Heading\nA note".encode())
        self.assertEqual([part["text"] for part in markdown["segments"]], ["# Heading", "A note"])

    def test_extracts_docx_paragraphs(self):
        output = io.BytesIO()
        document = Document()
        document.add_heading("Project plan", 1)
        document.add_paragraph("Review the requirements.")
        document.save(output)
        result = extract_document("docx", output.getvalue())
        self.assertIn("Project plan", " ".join(part["text"] for part in result["segments"]))

    def test_extracts_xlsx_rows(self):
        output = io.BytesIO()
        workbook = Workbook()
        workbook.active.title = "Budget"
        workbook.active.append(["Item", "Cost"])
        workbook.active.append(["Travel", 120])
        workbook.save(output)
        result = extract_document("xlsx", output.getvalue())
        self.assertEqual(result["segments"][1]["locator"], "Sheet Budget, row 2")
        self.assertIn("Travel | 120", result["segments"][1]["text"])

    def test_extracts_pptx_slide_text(self):
        output = io.BytesIO()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Quarterly review"
        presentation.save(output)
        result = extract_document("pptx", output.getvalue())
        self.assertEqual(result["segments"][0]["locator"], "Slide 1")

    def test_extracts_pdf_page_text(self):
        output = io.BytesIO()
        pdf = canvas.Canvas(output)
        pdf.drawString(72, 720, "Roadmap milestones")
        pdf.save()
        result = extract_document("pdf", output.getvalue())
        self.assertEqual(result["segments"][0]["locator"], "Page 1")

    def test_rejects_invalid_or_oversized_document(self):
        with self.assertRaisesRegex(ValueError, "valid PDF"):
            extract_document("pdf", b"not a pdf")
        with self.assertRaisesRegex(ValueError, "10 MB"):
            extract_document("txt", b"x" * (10 * 1024 * 1024 + 1))

    def test_rejects_zip_bomb_sized_package_before_office_parser(self):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"x" * (26 * 1024 * 1024))
        with self.assertRaisesRegex(ValueError, "extraction limits"):
            extract_document("docx", output.getvalue())


if __name__ == "__main__":
    unittest.main()
