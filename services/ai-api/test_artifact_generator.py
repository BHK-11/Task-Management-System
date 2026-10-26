import io
import unittest
from unittest.mock import patch

import artifact_generator
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from artifact_generator import generate_artifact


class ArtifactGeneratorTest(unittest.TestCase):
    def test_generates_and_reopens_editable_office_files_and_pdf(self):
        for kind in ("DOCX", "XLSX", "PPTX", "PDF"):
            with self.subTest(kind=kind):
                artifact = generate_artifact(kind, "Java revision plan", "## Core ideas\nReview classes and inheritance.\n- Practice examples")
                self.assertGreater(len(artifact.content), 100)
                if kind == "DOCX":
                    self.assertEqual(Document(io.BytesIO(artifact.content)).paragraphs[0].text, "Java revision plan")
                elif kind == "XLSX":
                    self.assertEqual(load_workbook(io.BytesIO(artifact.content), read_only=True).active["A1"].value, "Java revision plan")
                elif kind == "PPTX":
                    self.assertEqual(Presentation(io.BytesIO(artifact.content)).slides[0].shapes.title.text, "Java revision plan")
                else:
                    self.assertIn("Java revision plan", PdfReader(io.BytesIO(artifact.content)).pages[0].extract_text())

    def test_text_draft_is_utf8_and_not_a_fake_office_file(self):
        artifact = generate_artifact("TXT", "Teacher email draft", "Namaste, please find my assignment attached.")
        self.assertEqual(artifact.extension, "txt")
        self.assertEqual(artifact.content.decode("utf-8"), "Namaste, please find my assignment attached.")

    def test_xlsx_content_is_plain_text_not_an_executable_formula(self):
        artifact = generate_artifact("XLSX", "=HYPERLINK(\"https://example.invalid\")", "=WEBSERVICE(\"https://example.invalid\")")
        workbook = load_workbook(io.BytesIO(artifact.content), read_only=True, data_only=False)
        self.assertEqual(workbook.active["A1"].data_type, "s")
        self.assertEqual(workbook.active["A4"].data_type, "s")

    def test_rejects_missing_or_unbounded_content(self):
        with self.assertRaisesRegex(ValueError, "bounded"):
            generate_artifact("DOCX", "Title", " ")
        with self.assertRaisesRegex(ValueError, "bounded"):
            generate_artifact("DOCX", "Title", "x" * 12_001)

    def test_hindi_pdf_uses_installed_font_or_fails_clearly(self):
        with patch.dict("os.environ", {"TASKOS_PDF_FONT_PATH": ""}):
            if artifact_generator._system_devanagari_font():
                artifact = generate_artifact("PDF", "पढ़ाई", "यह एक उदाहरण है।")
                self.assertIn("पढ़ाई", PdfReader(io.BytesIO(artifact.content)).pages[0].extract_text())
            else:
                with self.assertRaisesRegex(ValueError, "Devanagari-capable"):
                    generate_artifact("PDF", "पढ़ाई", "यह एक उदाहरण है।")
