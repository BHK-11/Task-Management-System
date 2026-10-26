"""Deterministic, bounded generators for user-approved TASK OS deliverables."""

from __future__ import annotations

import html
import io
import os
import re
from dataclasses import dataclass

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches, Pt
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


MAX_CONTENT_CHARS = 12_000
MAX_OUTPUT_BYTES = 10 * 1024 * 1024


@dataclass(frozen=True)
class GeneratedArtifact:
    content: bytes
    media_type: str
    extension: str


def generate_artifact(kind: str, title: str, content: str) -> GeneratedArtifact:
    if not title.strip() or len(title) > 160 or not content.strip() or len(content) > MAX_CONTENT_CHARS:
        raise ValueError("A title and bounded, non-empty content are required")

    normalized = kind.upper()
    if normalized == "DOCX":
        result = _docx(title.strip(), content)
        media_type, extension = "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
    elif normalized == "XLSX":
        result = _xlsx(title.strip(), content)
        media_type, extension = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    elif normalized == "PPTX":
        result = _pptx(title.strip(), content)
        media_type, extension = "application/vnd.openxmlformats-officedocument.presentationml.presentation", "pptx"
    elif normalized == "PDF":
        result = _pdf(title.strip(), content)
        media_type, extension = "application/pdf", "pdf"
    elif normalized == "TXT":
        result = content.encode("utf-8")
        media_type, extension = "text/plain; charset=utf-8", "txt"
    else:
        raise ValueError("Unsupported deliverable format")

    if not result or len(result) > MAX_OUTPUT_BYTES:
        raise ValueError("The generated file is empty or exceeds the output limit")
    return GeneratedArtifact(result, media_type, extension)


def _lines(content: str) -> list[str]:
    return [line.rstrip() for line in content.replace("\r\n", "\n").split("\n")]


def _docx(title: str, content: str) -> bytes:
    document = Document()
    document.add_heading(title, level=0)
    for line in _lines(content):
        if not line.strip():
            continue
        if line.startswith("## "):
            document.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
        elif re.match(r"^[-*] ", line):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            document.add_paragraph(line)
    output = io.BytesIO()
    document.save(output)
    validated = Document(io.BytesIO(output.getvalue()))
    if not validated.paragraphs or validated.paragraphs[0].text != title:
        raise ValueError("DOCX verification failed")
    return output.getvalue()


def _xlsx(title: str, content: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Deliverable"
    sheet["A1"] = title
    sheet["A1"].data_type = "s"
    sheet["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="17243B")
    sheet["A3"] = "Content"
    sheet["A3"].font = Font(bold=True, color="FFFFFF")
    sheet["A3"].fill = PatternFill("solid", fgColor="315FE8")
    for index, line in enumerate(_lines(content), start=4):
        cell = sheet.cell(index, 1, line[:2000])
        cell.data_type = "s"
    sheet.column_dimensions["A"].width = 100
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:A{max(3, sheet.max_row)}"
    output = io.BytesIO()
    workbook.save(output)
    validated = load_workbook(io.BytesIO(output.getvalue()), read_only=True, data_only=False)
    if validated.active["A1"].value != title:
        raise ValueError("XLSX verification failed")
    return output.getvalue()


def _pptx(title: str, content: str) -> bytes:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = title
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = "Prepared with TASK OS"

    chunks: list[str] = []
    current: list[str] = []
    for line in _lines(content):
        if not line.strip():
            continue
        if line.startswith("## ") and current:
            chunks.append("\n".join(current))
            current = [line[3:].strip()]
        else:
            current.append(line[3:].strip() if line.startswith("## ") else line)
        if len("\n".join(current)) >= 900:
            chunks.append("\n".join(current))
            current = []
    if current:
        chunks.append("\n".join(current))
    chunks = chunks[:19]
    for index, text in enumerate(chunks, start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = text.splitlines()[0][:120]
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for paragraph_index, line in enumerate(text.splitlines()[1:] or [text]):
            paragraph = frame.paragraphs[0] if paragraph_index == 0 else frame.add_paragraph()
            paragraph.text = line[:600]
            paragraph.level = 0
            paragraph.font.size = Pt(18)
    output = io.BytesIO()
    presentation.save(output)
    validated = Presentation(io.BytesIO(output.getvalue()))
    if len(validated.slides) < 1 or validated.slides[0].shapes.title.text != title:
        raise ValueError("PPTX verification failed")
    return output.getvalue()


def _pdf(title: str, content: str) -> bytes:
    output = io.BytesIO()
    styles = getSampleStyleSheet()
    font_name = "Helvetica"
    font_path = os.getenv("TASKOS_PDF_FONT_PATH", "").strip() or _system_devanagari_font()
    has_devanagari = any("\u0900" <= char <= "\u097f" for char in title + content)
    if has_devanagari and not font_path:
        raise ValueError("A Devanagari-capable TASKOS_PDF_FONT_PATH or installed system font is required to verify Hindi PDF output")
    if has_devanagari:
        pdfmetrics.registerFont(TTFont("TaskOSUnicode", font_path))
        font_name = "TaskOSUnicode"
    styles.add(ParagraphStyle(name="TaskTitle", parent=styles["Title"], fontName=font_name, fontSize=20, leading=26))
    styles.add(ParagraphStyle(name="TaskBody", parent=styles["BodyText"], fontName=font_name, fontSize=10.5, leading=15, alignment=TA_LEFT, textColor=colors.HexColor("#17243b")))
    story = [Paragraph(html.escape(title), styles["TaskTitle"]), Spacer(1, 0.2 * inch)]
    for line in _lines(content):
        if not line.strip():
            story.append(Spacer(1, 0.08 * inch))
            continue
        text = line[3:].strip() if line.startswith("## ") else line
        story.append(Paragraph(html.escape(text), styles["TaskBody"]))
        story.append(Spacer(1, 0.08 * inch))
    SimpleDocTemplate(output, pagesize=letter, rightMargin=0.7 * inch, leftMargin=0.7 * inch,
                      topMargin=0.65 * inch, bottomMargin=0.65 * inch, title=title,
                      author="TASK OS").build(story)
    reader = PdfReader(io.BytesIO(output.getvalue()), strict=True)
    if not reader.pages or title not in (reader.pages[0].extract_text() or ""):
        raise ValueError("PDF verification failed")
    return output.getvalue()


def _system_devanagari_font() -> str:
    candidates = (
        r"C:\Windows\Fonts\mangal.ttf",
        "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",
        "/usr/share/fonts/truetype/lohit-devanagari/Lohit-Devanagari.ttf",
    )
    return next((path for path in candidates if os.path.isfile(path)), "")
